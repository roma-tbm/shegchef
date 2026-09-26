"""Оформление книги: цвета, рамки, ширины колонок, печать."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

INK = "1F2933"
MUTED = "7B8794"
LINE = "D6DBE1"
HEAD = "2E4756"
ACCENT = "1F6F5C"
WARN_BG = "FDF3E3"
WARN_FG = "9A5B0A"
DANGER_BG = "FBE9E7"
DANGER_FG = "A5320F"
OK_FG = "1F6F5C"
BAND = "F5F7F9"
ZEBRA = "FAFBFC"

THIN = Side(style="thin", color=LINE)

TITLE = Font(name="Calibri", size=16, bold=True, color=HEAD)
SUBTITLE = Font(name="Calibri", size=10, color=MUTED)
SECTION = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADFONT = Font(name="Calibri", size=10, bold=True, color=HEAD)
BODY = Font(name="Calibri", size=10, color=INK)
BOLD = Font(name="Calibri", size=10, bold=True, color=INK)
SMALL = Font(name="Calibri", size=9, color=MUTED)

WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
HEADALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws: Worksheet, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_title(ws: Worksheet, title: str, subtitle: str = "", row: int = 1) -> int:
    ws.cell(row=row, column=1, value=title).font = TITLE
    row += 1
    if subtitle:
        ws.cell(row=row, column=1, value=subtitle).font = SUBTITLE
        row += 1
    return row + 1


def section_bar(ws: Worksheet, row: int, title: str, span: int) -> int:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION
    cell.fill = PatternFill("solid", fgColor=HEAD)
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 22
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=HEAD)
    return row + 1


def table_header(ws: Worksheet, row: int, headers: tuple[str, ...]) -> int:
    for i, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = HEADFONT
        cell.fill = PatternFill("solid", fgColor=BAND)
        cell.alignment = HEADALIGN
        cell.border = Border(bottom=Side(style="medium", color=HEAD))
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return row + 1


def write_row(
    ws: Worksheet,
    row: int,
    values: tuple[object, ...],
    *,
    font: Font = BODY,
    zebra: bool = False,
    border: bool = True,
    fills: dict[int, str] | None = None,
    fonts: dict[int, Font] | None = None,
) -> int:
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.font = (fonts or {}).get(i, font)
        cell.alignment = WRAP
        if border:
            cell.border = Border(bottom=THIN)
        if zebra:
            cell.fill = PatternFill("solid", fgColor=ZEBRA)
        if fills and i in fills:
            cell.fill = PatternFill("solid", fgColor=fills[i])
    return row + 1


def merge_note(ws: Worksheet, row: int, text: str, span: int, *, kind: str = "info") -> int:
    style = {
        "info": (WARN_BG, WARN_FG),
        "danger": (DANGER_BG, DANGER_FG),
    }[kind]
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Calibri", size=9, color=style[1], bold=kind == "danger")
    cell.fill = PatternFill("solid", fgColor=style[0])
    cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = max(16, 14 * (1 + len(text) // (span * 14)))
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=style[0])
    return row + 1


def print_setup(ws: Worksheet, *, landscape: bool = True, fit_width: int = 1) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def fit_columns(ws: Worksheet, min_width: int = 8, max_width: int = 62) -> None:
    """Подгоняет ширину колонок по содержимому."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            longest = max((len(part) for part in text.split("\n")), default=0)
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(width + 2, min_width), max_width
        )
