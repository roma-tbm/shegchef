"""Чтение меню обратно из книги.

Пользователь правит лист «Конструктор меню» руками, поэтому чтение должно
быть терпимым к мелочам: заголовок не в первой строке, блюдо набрано с
пробелами, порции — текстом, а половина строк может быть пустой.

Лист «План меню» читается так же: из него берётся смена за конкретный день
недели, когда шеф запускает сборку с ключом --week.
"""

from __future__ import annotations

from collections.abc import Iterator, Mapping
from dataclasses import dataclass, replace
from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook

from kitchen.excel.folders import MEAL_ORDER
from kitchen.menu_week import WEEKDAYS, normalize_weekday
from kitchen.models import MenuLine, PlanLine

MENU_SHEET = "Конструктор меню"
PLAN_SHEET = "План меню"
HEADER_WINDOW = 30
"""Сколько первых строк просматриваем в поисках заголовка."""

DEFAULT_MEAL = "Обед"
DEFAULT_SERVE = time(12, 0)

# Синонимы заголовков: пользователь мог переименовать колонку.
DISH_ALIASES = ("Блюдо", "Наименование", "Блюдо/наименование")
PORTIONS_ALIASES = ("Порций", "Порции", "Порций, чел")
SERVE_ALIASES = ("Выдача", "Время выдачи", "К выдаче")
MEAL_ALIASES = ("Приём пищи", "Прием пищи", "Тип приёма пищи", "Приём")
WEEKDAY_ALIASES = ("День недели", "День", "Неделя")
NOTE_ALIASES = ("Примечание", "Примечания", "Комментарий")


def _index(cols: dict[str, int], names: tuple[str, ...]) -> int | None:
    """Номер колонки по первому подходящему заголовку-синониму."""
    for name in names:
        if name in cols:
            return cols[name]
    return None


def _parse_time(value, default: time) -> time:
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%H:%M", "%H.%M", "%H:%M:%S"):
            try:
                from datetime import datetime

                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return default


@dataclass(frozen=True)
class Block:
    """Кусок листа «Конструктор меню», который можно читать как таблицу.

    Лист устроен окнами: заголовок «Завтрак», своя шапка, строки позиций,
    потом «Обед» и так далее. Старые книги плоские — там одна таблица без
    заголовков-папок, и она читается тем же кодом, просто блок один.
    """

    meal: str
    header_row: int
    cols: dict[str, int]
    first: int
    last: int

    @property
    def dish(self) -> int | None:
        """Номер колонки «Блюдо», считая с единицы, — как в openpyxl."""
        return _index(self.cols, DISH_ALIASES)

    @property
    def portions(self) -> int | None:
        return _index(self.cols, PORTIONS_ALIASES)

    @property
    def serve(self) -> int | None:
        return _index(self.cols, SERVE_ALIASES)

    @property
    def note(self) -> int | None:
        return _index(self.cols, NOTE_ALIASES)


def _cells(row) -> list:
    return list(row)


def _is_bar(row) -> bool:
    """Заголовок-полоса: текст в первой колонке и больше ничего в строке."""
    cells = _cells(row)
    if not cells or not cells[0].value:
        return False
    return all(not c.value for c in cells[1:])


def _is_header(labels: set[str]) -> bool:
    return (
        any(n in labels for n in DISH_ALIASES)
        and any(n in labels for n in PORTIONS_ALIASES)
    )


def _columns(row) -> dict[str, int]:
    """Заголовки строки → номера колонок с единицы, как их адресует openpyxl."""
    labels = [str(c.value).strip() if c.value else "" for c in row]
    return {name: i + 1 for i, name in enumerate(labels) if name}


def _meal_candidates(recipes: Mapping[str, object] | None) -> set[str]:
    """Названия приёмов пищи, которые считаем заголовком окна.

    Кроме канонического списка берём приёмы из каталога: шеф мог завести
    «Второй завтрак», и такое окно тоже должно читаться.
    """
    names = set(MEAL_ORDER)
    for recipe in (recipes or {}).values():
        meal = getattr(recipe, "meal", None)
        if meal:
            names.add(str(meal))
    return names


def _match_meal(text: str, candidates: set[str]) -> str:
    needle = text.strip().casefold()
    for meal in candidates:
        if meal.casefold() == needle:
            return meal
    return ""


def iter_blocks(ws, recipes: Mapping[str, object] | None = None) -> Iterator[Block]:
    """Идёт по листу и отдаёт окна меню — то, что действительно заполнено.

    Панель поиска и шапка листа пропускаются: блок начинается только после
    строки заголовков. Поэтому счётчики формул нельзя принять за блюдо,
    даже если Excel их уже посчитал.
    """
    candidates = _meal_candidates(recipes)
    header = 0
    cols: dict[str, int] = {}
    first = 0
    meal = ""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        index = row[0].row
        if _is_bar(row):
            if header:
                yield Block(meal, header, cols, first, _last_row(ws, cols, first, index - 1))
            found = _match_meal(str(row[0].value or ""), candidates)
            header, cols, first = 0, {}, index + 1
            meal = found or ""
            continue
        labels = {str(c.value).strip() for c in row if c.value is not None}
        if _is_header(labels):
            if header:
                yield Block(meal, header, cols, first, _last_row(ws, cols, first, index - 1))
            header, cols, first = index, _columns(row), index + 1
    if header:
        yield Block(meal, header, cols, first, _last_row(ws, cols, first, ws.max_row))


def _last_row(ws, cols: Mapping[str, int], first: int, limit: int) -> int:
    """Последняя строка окна: пустая строка-разделитель между окнами не наша.

    Окна разделены пустой строкой, и если её не отбросить, подбор блюд впишет
    туда позицию — без выпадающего списка и рамки, а шеф её не найдёт.
    """
    last = min(limit, ws.max_row)
    while last > first and not _has_values(ws, cols, last):
        last -= 1
    return last


def _has_values(ws, cols: Mapping[str, int], row: int) -> bool:
    return any(
        ws.cell(row=row, column=col).value not in (None, "")
        for col in cols.values()
    )


def read_menu(path: Path, recipes: dict[str, object] | None = None) -> tuple[MenuLine, ...]:
    """Читает меню из книги. Пустой лист — это пустое меню, а не ошибка."""
    wb = load_workbook(path, data_only=True)
    try:
        if MENU_SHEET not in wb.sheetnames:
            return ()
        return _parse_sheet(wb[MENU_SHEET], recipes or {})
    finally:
        wb.close()


def _parse_sheet(ws, recipes: dict[str, object]) -> tuple[MenuLine, ...]:
    rows: list[MenuLine] = []
    index = _dish_index(recipes)
    merged: dict[tuple[str, str], MenuLine] = {}
    for block in iter_blocks(ws, recipes):
        if block.dish is None or block.portions is None:
            continue
        serve_col = _index(block.cols, SERVE_ALIASES)
        meal_col = _index(block.cols, MEAL_ALIASES)
        for row in ws.iter_rows(min_row=block.first, max_row=block.last):
            values = [c.value for c in row]
            dish = values[block.dish - 1]
            if dish is None or not str(dish).strip():
                continue
            typed = str(dish).strip()

            portions = _parse_int(values[block.portions - 1])
            if portions is None or portions <= 0:
                # Ноль порций — это «блюдо убрали из меню», а не ошибка ввода.
                continue

            name = _match_recipe(typed, index) if index else typed
            if name is None:
                _warn(f"пропущено «{typed}» — нет такого блюда в ТТК, "
                      "впишите название как в «ТТК»")
                continue

            meal = block.meal
            if not meal and meal_col is not None and values[meal_col - 1]:
                meal = str(values[meal_col - 1]).strip()
            if meal not in MEAL_ORDER:
                meal = block.meal or DEFAULT_MEAL

            serve = (
                _parse_time(values[serve_col - 1], DEFAULT_SERVE)
                if serve_col is not None
                else DEFAULT_SERVE
            )
            already = merged.get((meal, name))
            if already is not None:
                # Одно блюдо в двух строках окна — это удвоенная норма продуктов.
                merged[(meal, name)] = replace(
                    already, portions=already.portions + portions
                )
                _warn(f"«{name}» в окне «{meal}» записан дважды — порции сложены")
                continue
            line = MenuLine(name, portions, serve, meal, date.today())
            merged[(meal, name)] = line
            rows.append(line)
    return tuple(rows)


def _norm(text: str) -> str:
    return " ".join(str(text).split()).casefold()


def _dish_index(recipes: Mapping[str, object]) -> dict[str, str]:
    """Нормализованное имя → точное название в ТТК.

    Если два блюда отличаются только регистром или пробелами, ключ убирается:
    иначе одно из них потерялось бы молча, и впиши шеф любое из двух — меню
    получило бы не то блюдо.
    """
    index: dict[str, str] = {}
    for name in recipes:
        key = _norm(name)
        if key in index:
            index[key] = ""
        else:
            index[key] = name
    return {key: name for key, name in index.items() if name}


def _match_recipe(typed: str, index: Mapping[str, str]) -> str | None:
    """Ищет блюдо в ТТК так, чтобы шефу не приходилось выбирать из списка.

    Выпадающий список работает не во всех программах — Numbers его не
    показывает, — и блюдо часто вписывают руками. Регистр и лишние пробелы
    не должны быть причиной потерять строку меню. Начало названия тоже
    годится, но только если оно указывает на одно блюдо: «Салат» — это
    четыре блюда, а вот «Салат «Цезарь»» — одно.
    """
    key = _norm(typed)
    if key in index:
        return index[key]
    matches = [name for norm, name in index.items() if norm.startswith(key)]
    if len(matches) == 1:
        return matches[0]
    return None


def _find_header(ws, required=None) -> tuple[int | None, dict[str, int]]:
    """Ищет строку заголовков: «Блюдо» + «Порций» + необязательный третий признак."""
    groups = [DISH_ALIASES, PORTIONS_ALIASES]
    if required:
        groups.append(required)
    for row in ws.iter_rows(min_row=1, max_row=HEADER_WINDOW):
        labels = [str(c.value).strip() if c.value else "" for c in row]
        if all(any(n in labels for n in group) for group in groups):
            return row[0].row, {
                name: labels.index(name) + 1 for name in labels if name
            }
    return None, {}


def _parse_int(value) -> int | None:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def _warn(message: str) -> None:
    import sys

    print(f"  ! {message}", file=sys.stderr)


# ---------------------------------------------------------------------------
# План меню на неделю
# ---------------------------------------------------------------------------


def plan_for(plan: tuple[PlanLine, ...], weekday: str) -> tuple[MenuLine, ...]:
    """Строки плана за день недели — готовые строки меню.

    Нужен Excel-слою, потому что «План меню» шеф тоже правит руками: план
    может прийти из книги и не совпасть с kitchen/menu_week.py.
    """
    name = normalize_weekday(weekday)
    if name is None:
        raise ValueError(
            f"Неизвестный день недели: «{weekday}». Ожидается один из: "
            f"{', '.join(WEEKDAYS)} (или сокращение: пн, вт, ср, чт, пт)."
        )
    return tuple(
        MenuLine(line.recipe, line.portions, line.serve_at, line.meal, date.today())
        for line in plan
        if line.weekday == name
    )


def read_plan(
    path: Path, recipes: dict[str, object] | None = None
) -> tuple[PlanLine, ...]:
    """Читает недельный план из книги. Пустой лист — это пустой план."""
    wb = load_workbook(path, data_only=True)
    try:
        if PLAN_SHEET not in wb.sheetnames:
            return ()
        return _parse_plan(wb[PLAN_SHEET], recipes or {})
    finally:
        wb.close()


def _parse_plan(ws, recipes: dict[str, object]) -> tuple[PlanLine, ...]:
    header_row, cols = _find_header(ws, WEEKDAY_ALIASES)
    if header_row is None:
        return ()

    dish_col = _index(cols, DISH_ALIASES)
    portions_col = _index(cols, PORTIONS_ALIASES)
    if dish_col is None or portions_col is None:
        return ()
    day_col = _index(cols, WEEKDAY_ALIASES)
    serve_col = _index(cols, SERVE_ALIASES)
    meal_col = _index(cols, MEAL_ALIASES)
    note_col = _index(cols, NOTE_ALIASES)

    rows: list[PlanLine] = []
    day = ""
    for row in ws.iter_rows(min_row=header_row + 1):
        values = [c.value for c in row]
        if day_col is not None and values[day_col - 1]:
            day = normalize_weekday(str(values[day_col - 1])) or ""
        dish = values[dish_col - 1]
        if not day or dish is None or not str(dish).strip():
            continue
        dish = str(dish).strip()

        if recipes and dish not in recipes:
            _warn(f"в плане пропущено «{dish}» — нет такого блюда в ТТК")
            continue

        portions = _parse_int(values[portions_col - 1])
        if portions is None or portions <= 0:
            continue

        meal = str(values[meal_col - 1]).strip() if meal_col and values[meal_col - 1] else ""
        if meal not in MEAL_ORDER:
            meal = DEFAULT_MEAL
        serve = (
            _parse_time(values[serve_col - 1], DEFAULT_SERVE)
            if serve_col
            else DEFAULT_SERVE
        )
        note = str(values[note_col - 1]).strip() if note_col and values[note_col - 1] else ""
        rows.append(PlanLine(day, meal, dish, portions, serve, note))
    return tuple(rows)


__all__ = [
    "MENU_SHEET",
    "PLAN_SHEET",
    "WEEKDAYS",
    "normalize_weekday",
    "plan_for",
    "read_menu",
    "read_plan",
]
