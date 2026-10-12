"""Регрессионные тесты дефектов четвёртого ревью Stage 3 (V4-01, V4-02).

    V4-01  книга Excel теряла хозяйственные работы и подготовку следующего дня
    V4-02  сохранённый расход сырья терял единицу измерения

Тесты независимы от прежних проверок: они сравнивают рабочие представления с
полным авторитетным планом (в обе стороны) и проверяют единицу расхода после
перезапуска, смены и очистки меню.
"""

from __future__ import annotations

from datetime import date, time
from pathlib import Path

import pytest

from kitchen.duties import DUTIES_BY_KEY, DutySettings
from kitchen.models import (
    Ingredient,
    KitchenData,
    MenuLine,
    Product,
    Recipe,
    TaskTemplate,
)
from kitchen.staff import ROLE_CLEANER, ROLE_COOK, Employee
from kitchen.web import report_store
from kitchen.web.report_store import (
    SaturdayReport,
    SemifinishedMarking,
    SemifinishedOutput,
)
from kitchen.web.shift_plan import build_scenario, saturday_report_form

MONDAY = date(2026, 10, 5)
TUESDAY = date(2026, 10, 6)
SATURDAY = date(2026, 10, 10)

POTATO = Product(name="Картофель", unit="кг", category="Овощи", storage="Овощная")

STAGE_COOK = "Приготовление"
STAGE_PREP = "Заготовка"


def _recipe(name: str, *, prep_min: int = 0, cook_min: int = 10) -> Recipe:
    return Recipe(
        name=name,
        meal="Обед",
        category="Обед",
        subcategory="Второе",
        shop="Горячий цех",
        prep_min=prep_min,
        cook_min=cook_min,
        ingredients=(Ingredient("Картофель", 200.0),),
    )


def _cook_template(name: str, *, minutes: int = 10) -> TaskTemplate:
    return TaskTemplate(
        recipe=name,
        order=1,
        stage=STAGE_COOK,
        role=ROLE_COOK,
        operation=f"Cook {name}",
        duration_min=minutes,
        products="Картофель",
        meal="Обед",
    )


def _prep_template(name: str, *, minutes: int = 10) -> TaskTemplate:
    return TaskTemplate(
        recipe=name,
        order=1,
        stage=STAGE_PREP,
        role=ROLE_COOK,
        operation=f"Prep {name}",
        duration_min=minutes,
        products="Картофель",
        meal="Обед",
    )


def _line(name: str, day: date, *, portions: int = 20, serve_at: time = time(12, 0)) -> MenuLine:
    return MenuLine(recipe=name, portions=portions, serve_at=serve_at, meal="Обед", day=day)


def _employee(name: str, role: str, start: time, end: time, order: int) -> Employee:
    return Employee(
        name=name, role=role, shift_start=start, shift_end=end, breaks=(), order=order
    )


def _v4_data() -> KitchenData:
    return KitchenData(
        products={"Картофель": POTATO},
        recipes={"Cook A": _recipe("Cook A"), "Cook B": _recipe("Cook B")},
        tasks=(_cook_template("Cook A"), _prep_template("Cook B")),
        menu=(_line("Cook A", MONDAY),),
    )


#: Только подготовка следующего дня и финальная мойка: остальное отключено.
def _v4_duties() -> DutySettings:
    keep = {"next-day-prep", "final-kitchen-clean"}
    return DutySettings(
        disabled=frozenset(DUTIES_BY_KEY) - keep,
        duration_overrides={"next-day-prep": 0, "final-kitchen-clean": 10},
    )


def _v4_scenario():
    data = _v4_data()
    return build_scenario(
        data,
        MONDAY,
        menu=(_line("Cook A", MONDAY),),
        next_menu=(_line("Cook B", TUESDAY),),
        staff=(
            _employee("Cook", ROLE_COOK, time(7, 0), time(20, 0), 1),
            _employee("Cleaner", ROLE_CLEANER, time(8, 0), time(17, 0), 2),
        ),
        duty_settings=_v4_duties(),
    ), data


def _replace_menu(data: KitchenData, menu: tuple[MenuLine, ...]) -> KitchenData:
    from dataclasses import replace

    return replace(data, menu=menu)


# ---------------------------------------------------------------------------
# V4-01 — рабочие представления отражают полный авторитетный план
# ---------------------------------------------------------------------------


def test_plan_contains_household_work_and_next_day_prep() -> None:
    """План действительно содержит хозяйственную работу и подготовку."""
    scenario, _ = _v4_scenario()

    operations = {i.operation for i in scenario.plan.items}
    assert any("на следующий день" in op for op in operations)
    assert any("Финальная мойка" in op for op in operations)
    # Подготовка стоит в своём окне 14:00, финальная мойка — 16:30–16:40.
    prep = next(i for i in scenario.plan.items if "на следующий день" in i.operation)
    clean = next(i for i in scenario.plan.items if "Финальная мойка" in i.operation)
    assert prep.start.time() == time(14, 0)
    assert (clean.start.time(), clean.end.time()) == (time(16, 30), time(16, 40))
    assert clean.assignee == "Cleaner"


def test_sheet_keeps_every_plan_item_and_adds_nothing_foreign() -> None:
    """Полнота в обе стороны: ни одна работа не потеряна и нет чужих строк."""
    from kitchen.core import build_shift_sheet

    scenario, data = _v4_scenario()
    menu = (tuple(data.menu)[0],)
    sheet = build_shift_sheet(_replace_menu(data, menu), plan=scenario.plan)

    expected = {i.item_id: (i.start, i.end) for i in scenario.plan.items}
    actual = {t.item_id: (t.start, t.end) for t in sheet.tasks}
    assert set(actual) == set(expected), {
        "потеряно": sorted(set(expected) - set(actual)),
        "лишнее": sorted(set(actual) - set(expected)),
    }
    for item_id, span in expected.items():
        assert actual[item_id] == span, item_id


def test_sheet_shift_bounds_follow_the_full_plan() -> None:
    """Границы смены — по последней работе плана, а не по старому таймингу."""
    from kitchen.core import build_shift_sheet

    scenario, data = _v4_scenario()
    menu = (tuple(data.menu)[0],)
    sheet = build_shift_sheet(_replace_menu(data, menu), plan=scenario.plan)

    assert sheet.shift_start.time() == time(11, 50)
    assert sheet.shift_end.time() == time(16, 40)


def test_sheet_carries_assignees_statuses_and_comments() -> None:
    """Исполнитель, статус и комментарий переносятся в строку книги."""
    from dataclasses import replace as dc_replace

    from kitchen.core import build_shift_sheet

    scenario, data = _v4_scenario()
    # Отмечаем комментарий человеком по устойчивому item_id одной работы.
    target = next(i for i in scenario.plan.items if "на следующий день" in i.operation)
    scenario = dc_replace(
        scenario,
        plan=dc_replace(
            scenario.plan,
            items=tuple(
                dc_replace(
                    i,
                    comment="сдвинули из-за поставки",
                    status="IN_PROGRESS",
                )
                if i.item_id == target.item_id
                else i
                for i in scenario.plan.items
            ),
        ),
    )
    menu = (tuple(data.menu)[0],)
    sheet = build_shift_sheet(_replace_menu(data, menu), plan=scenario.plan)

    line = next(t for t in sheet.tasks if t.item_id == target.item_id)
    assert line.assignee == "Cook"
    assert line.status == "IN_PROGRESS"
    assert line.comment == "сдвинули из-за поставки"


def test_workbook_shows_duties_and_next_day_prep(tmp_path) -> None:
    """Реальная книга содержит финальную мойку и подготовку следующего дня."""
    from openpyxl import load_workbook

    from kitchen.core import build_shift_sheet
    from kitchen.excel.builder import build_workbook

    scenario, data = _v4_scenario()
    menu = (tuple(data.menu)[0],)
    book_data = _replace_menu(data, menu)
    sheet = build_shift_sheet(book_data, plan=scenario.plan)

    path = tmp_path / "Кухня.xlsx"
    build_workbook(book_data, sheet, path)

    wb = load_workbook(path)
    text = "\n".join(
        str(cell.value)
        for ws in (wb["Лист смены"], wb["Карта задач"], wb["Тайминг"])
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "Финальная мойка" in text
    assert "на следующий день" in text
    assert "16:40" in text, "граница смены в книге взята не из полного плана"
    assert "Cleaner" in text, "исполнитель хозяйственной работы потерян"


def test_plan_control_point_reaches_the_sheet_with_executor() -> None:
    """Контрольная точка плана попадает и в строки, и в блок контроля."""
    from kitchen.core import build_shift_sheet

    data = KitchenData(
        products={"Картофель": POTATO},
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(
            _cook_template("Cook A"),
            TaskTemplate(
                recipe="Cook A",
                order=2,
                stage="Контроль",
                role=ROLE_COOK,
                operation="Проверка Cook A",
                duration_min=5,
                meal="Обед",
            ),
        ),
        menu=(_line("Cook A", MONDAY),),
    )
    scenario = build_scenario(
        data,
        MONDAY,
        menu=(_line("Cook A", MONDAY),),
        staff=(_employee("Cook", ROLE_COOK, time(7, 0), time(20, 0), 1),),
        duty_settings=DutySettings(disabled=frozenset(DUTIES_BY_KEY)),
    )
    book_data = _replace_menu(data, tuple(data.menu))
    sheet = build_shift_sheet(book_data, plan=scenario.plan)

    line = next(t for t in sheet.tasks if t.operation == "Проверка Cook A")
    assert line.assignee == "Cook"
    control = next(c for c in sheet.controls if c.operation == "Проверка Cook A")
    assert control.assignee == "Cook"


def test_workbook_keeps_expected_sheets(tmp_path) -> None:
    """Полный план не меняет набор и порядок листов книги."""
    from openpyxl import load_workbook

    from kitchen.core import build_shift_sheet
    from kitchen.excel.builder import build_workbook

    scenario, data = _v4_scenario()
    menu = (tuple(data.menu)[0],)
    book_data = _replace_menu(data, menu)
    path = tmp_path / "Кухня.xlsx"
    build_workbook(book_data, build_shift_sheet(book_data, plan=scenario.plan), path)

    wb = load_workbook(path)
    for sheet in ("Лист смены", "Меню", "Продукты", "Карта задач", "Тайминг"):
        assert sheet in wb.sheetnames, wb.sheetnames


# ---------------------------------------------------------------------------
# V4-02 — единица фактического расхода сохраняется вместе с количеством
# ---------------------------------------------------------------------------


def _report_with_unit(day: date) -> SaturdayReport:
    return SaturdayReport(
        day=day,
        consumption={"Картофель": 5.0},
        consumption_units={"Картофель": "кг"},
        outputs={"Бульон": SemifinishedOutput(qty=4.0, unit="л")},
        markings={
            "Бульон": SemifinishedMarking(
                name="Бульон", made=day, expires=date(2026, 10, 13)
            )
        },
    )


def _facts(report: SaturdayReport | None) -> dict:
    if report is None:
        return {}
    return {
        "actuals": dict(report.consumption),
        "actual_units": dict(report.consumption_units),
        "outputs": dict(report.outputs),
        "markings": dict(report.markings),
    }


def test_consumption_unit_round_trips_through_storage(tmp_path) -> None:
    """Единица расхода переживает сохранение и повторное чтение."""
    path = tmp_path / "reports.json"
    report_store.save_report(SATURDAY, _report_with_unit(SATURDAY), path)

    loaded = report_store.load_report(SATURDAY, path)

    assert loaded is not None
    assert loaded.consumption == {"Картофель": 5.0}
    assert loaded.consumption_units == {"Картофель": "кг"}


def test_export_keeps_consumption_and_unit_after_menu_is_cleared(tmp_path) -> None:
    """Пустое меню не лишает сохранённый расход его единицы."""
    path = tmp_path / "reports.json"
    report_store.save_report(SATURDAY, _report_with_unit(SATURDAY), path)
    saved = report_store.load_report(SATURDAY, path)

    scenario = build_scenario(KitchenData(), SATURDAY, menu=())
    form = saturday_report_form(scenario, **_facts(saved))

    row = next(r for r in form.rows if r[0] == "Картофель")
    # Продукт, количество и исходная единица — все три вместе, а не прочерк.
    assert row[3] == "5"
    assert row[4] == "кг"
    assert "кг" in form.to_text()


def test_consumption_unit_is_not_taken_from_output(tmp_path) -> None:
    """Единица выхода не подменяет единицу расхода при общем ключе."""
    report = SaturdayReport(
        day=SATURDAY,
        consumption={"Бульон": 5.0},
        consumption_units={},  # единица расхода неизвестна
        outputs={"Бульон": SemifinishedOutput(qty=4.0, unit="кг")},
        markings={
            "Бульон": SemifinishedMarking(
                name="Бульон", made=SATURDAY, expires=date(2026, 10, 13)
            )
        },
    )
    scenario = build_scenario(KitchenData(), SATURDAY, menu=())

    form = saturday_report_form(scenario, **_facts(report))

    row = next(r for r in form.rows if r[0] == "Бульон")
    assert row[4] == "—", "единица расхода не выводится из единицы выхода"
    assert any("Единица расхода не указана" in n for n in form.notes)


def test_legacy_json_without_units_reads_compatibly(tmp_path) -> None:
    """Старый JSON без единиц читается, неполнота показывается, а не гадается."""
    path = tmp_path / "reports.json"
    path.write_text(
        '{"version": 1, "days": {"2026-10-10": {'
        '"consumption": {"Картофель": 5.0}, '
        '"outputs": {}, "markings": {}}}}',
        encoding="utf-8",
    )

    loaded = report_store.load_report(SATURDAY, path)

    assert loaded is not None
    assert loaded.consumption == {"Картофель": 5.0}
    assert loaded.consumption_units == {}

    scenario = build_scenario(KitchenData(), SATURDAY, menu=())
    form = saturday_report_form(scenario, **_facts(loaded))
    row = next(r for r in form.rows if r[0] == "Картофель")
    assert row[4] == "—"
    assert any("Единица расхода не указана" in n for n in form.notes)


def test_cli_printout_keeps_the_consumption_unit(tmp_path) -> None:
    """Реальный CLI-экспорт печатает единицу расхода вместе с количеством."""
    import build as cli

    out = tmp_path / "Кухня.xlsx"
    report_store.save_report(
        SATURDAY, _report_with_unit(SATURDAY), tmp_path / "reports.json"
    )

    assert cli.main(["--out", str(out), "--day", "10.10.2026", "--menu", "50",
                     "--reports", str(tmp_path / "reports.json"), "--no-open"]) == 0

    forms = list((tmp_path / cli.PRINT_DIR_NAME).glob("Субботний отчёт*.txt"))
    assert forms, "субботний отчёт должен быть выгружен"
    text = forms[0].read_text(encoding="utf-8")
    assert "Картофель" in text
    assert "5" in text
    # Единица выхода здесь «л», поэтому «кг» может прийти только из расхода.
    assert "кг" in text
    assert "л" in text


def test_app_export_uses_saved_consumption_unit(tmp_path, monkeypatch) -> None:
    """Экспорт интерфейса берёт единицу расхода из сохранённого отчёта."""
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    from kitchen.web import menu_store
    from kitchen.web.dataset import Dataset

    app_path = Path(__file__).resolve().parent.parent / "app.py"
    monkeypatch.setenv("SHEGCHEF_DATA_DIR", str(tmp_path))
    store = Dataset.in_dir(tmp_path).menu
    menu_store.save_days({SATURDAY.isoformat(): ()}, store)
    report_store.save_report(
        SATURDAY, _report_with_unit(SATURDAY), Dataset.in_dir(tmp_path).reports
    )

    tester = AppTest.from_file(app_path, default_timeout=120).run()
    assert not tester.exception, [e.value for e in tester.exception]
    tester.session_state["day"] = SATURDAY
    tester.session_state["menu"] = []
    tester.run()
    assert not tester.exception, [e.value for e in tester.exception]

    saved = report_store.load_report(SATURDAY, tmp_path / "reports.json")
    assert saved is not None and saved.consumption_units == {"Картофель": "кг"}
    labels = [b.label for b in tester.tabs[5].download_button]
    assert any("Субботний отчёт" in label for label in labels), labels
