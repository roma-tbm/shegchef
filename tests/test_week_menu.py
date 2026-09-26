"""Проверки недельного меню: план разворачивается в смену, ссылки рабочие.

Эти тесты повторяют второй рабочий сценарий пользователя — «а что бы я
приготовил в среду?» — и проверяют, что книга осталась цельной после того,
как каталог вырос с 5 блюд до 29.
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date, time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kitchen.core import build_shift_sheet
from conftest import column_of, header_of, row_of

from kitchen.excel import (
    RECIPE_REF_SHEET,
    build_workbook,
    normalize_weekday,
    plan_for,
    read_plan,
)
from kitchen.menu_week import (
    PLAN_PORTIONS,
    REUSED_RECIPES,
    WEEK_PLAN,
    WEEK_RECIPES,
    WEEK_TASKS,
    WEEKDAYS,
)
from kitchen.models import Ingredient, MenuLine, PlanLine, Recipe
from kitchen.seed import CATALOG_PRODUCTS, CATALOG_RECIPES, demo_data

DAY = date(2026, 9, 26)


@pytest.fixture()
def book_path(tmp_path: Path) -> Path:
    data = demo_data(DAY)
    path = tmp_path / "Кухня.xlsx"
    return build_workbook(data, build_shift_sheet(data), path)


# ---------------------------------------------------------------------------
# Данные недели
# ---------------------------------------------------------------------------


def test_план_покрывает_пять_рабочих_дней():
    assert {p.weekday for p in WEEK_PLAN} == set(WEEKDAYS)


def test_на_каждый_день_есть_завтрак_и_обед():
    """Пустой день недели — это не «план», а опечатка."""
    by_day: dict[str, set[str]] = {}
    for line in WEEK_PLAN:
        by_day.setdefault(line.weekday, set()).add(line.meal)
    for day, meals in by_day.items():
        assert "Завтрак" in meals, f"{day}: нет завтрака"
        assert "Обед" in meals, f"{day}: нет обеда"


def test_каждая_позиция_плана_есть_в_ттк():
    known = set(CATALOG_RECIPES)
    unknown = {p.recipe for p in WEEK_PLAN} - known
    assert not unknown, f"в плане есть блюда без карточки: {sorted(unknown)}"


def test_каждый_ингредиент_есть_в_каталоге_продуктов():
    known = {p.name for p in CATALOG_PRODUCTS}
    missing = {
        (r.name, i.product)
        for r in WEEK_RECIPES
        for i in r.ingredients
        if i.product not in known
    }
    assert not missing, f"нет таких продуктов: {sorted(missing)}"


def test_у_каждого_нового_блюда_есть_операции():
    """Карточка без операций попадёт в меню, но не даст ни одной задачи."""
    covered = {t.recipe for t in WEEK_TASKS}
    orphans = {r.name for r in WEEK_RECIPES} - covered
    assert not orphans, f"без операций: {sorted(orphans)}"


def test_строки_операций_не_содержат_запятых():
    """Запятая в «Продуктах» разделяет список, поэтому название без запятой."""
    for tpl in WEEK_TASKS:
        if tpl.products:
            for name in tpl.products.split(", "):
                assert name in {p.name for p in CATALOG_PRODUCTS}, (
                    f"{tpl.recipe}: «{tpl.operation}» — неизвестный продукт {name!r}"
                )


def test_переиспользованные_блюда_не_заведены_дважды():
    """Три позиции взяты из старого ТТК — значит, в каталоге они по одному разу."""
    for name in REUSED_RECIPES:
        assert name in CATALOG_RECIPES
        assert name not in {r.name for r in WEEK_RECIPES}


def test_каталог_вырос_но_демо_меню_не_тронуто():
    assert len(CATALOG_RECIPES) == 29
    assert len(demo_data(DAY).menu) == 5
    assert len(demo_data(DAY).plan) == 27


def test_в_плане_нет_битых_символов():
    for line in WEEK_PLAN:
        for text in (line.recipe, line.note):
            assert "�" not in text, f"битый символ в {text!r}"


# ---------------------------------------------------------------------------
# День недели
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("пн", "Понедельник"),
        ("ПН", "Понедельник"),
        ("понедельник", "Понедельник"),
        ("Понедельник", "Понедельник"),
        ("  вт  ", "Вторник"),
        ("2", "Вторник"),
        ("чтв", "Четверг"),
        ("пт", "Пятница"),
    ],
)
def test_день_недели_понимается_во_всех_форматах(value: str, expected: str):
    assert normalize_weekday(value) == expected


@pytest.mark.parametrize("value", ["", "суббота", "6", "0", "понедельникн"])
def test_неизвестный_день_недели(value: str):
    assert normalize_weekday(value) is None
    with pytest.raises(ValueError, match="день недели"):
        plan_for(WEEK_PLAN, value)


def test_план_даёт_меню_за_день():
    rows = plan_for(WEEK_PLAN, "ср")
    assert rows
    assert all(m.portions == PLAN_PORTIONS for m in rows)
    assert {m.meal for m in rows} == {"Завтрак", "Обед"}


def test_дни_недели_не_пересекаются():
    """Иначе один и тот же продукт уехал бы в закупку дважды за день."""
    seen: set[tuple[str, str]] = set()
    for line in WEEK_PLAN:
        key = (line.weekday, line.recipe)
        assert key not in seen, f"{line.weekday}: {line.recipe} повторяется"
        seen.add(key)


# ---------------------------------------------------------------------------
# Книга
# ---------------------------------------------------------------------------


def test_план_читается_из_книги(book_path: Path):
    from_book = read_plan(book_path, CATALOG_RECIPES)
    assert len(from_book) == len(WEEK_PLAN)
    assert {p.weekday for p in from_book} == set(WEEKDAYS)
    assert {p.recipe for p in from_book} == {p.recipe for p in WEEK_PLAN}


def test_пустой_план_не_ломает_сборку(tmp_path: Path):
    data = replace(demo_data(DAY), menu=demo_data(DAY).menu, plan=())
    path = build_workbook(data, build_shift_sheet(data), tmp_path / "без_плана.xlsx")
    assert read_plan(path, CATALOG_RECIPES) == ()
    assert "План меню" in load_workbook(path).sheetnames


def test_план_меню_в_книге_группируется_по_дням(book_path: Path):
    """Каждый день — свой заголовок, а не 27 строк в одну колонку."""
    ws = load_workbook(book_path)["План меню"]
    days = [
        ws.cell(row=r, column=1).value
        for r in range(1, ws.max_row + 1)
        if ws.cell(row=r, column=1).value in WEEKDAYS
    ]
    assert days == list(WEEKDAYS)


def test_справочник_блюд_скрыт_и_полон(book_path: Path):
    wb = load_workbook(book_path)
    ws = wb[RECIPE_REF_SHEET]
    assert ws.sheet_state == "hidden"
    names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert set(names) == set(CATALOG_RECIPES)


def test_выпадающий_список_ссылается_на_справочник(book_path: Path):
    """29 блюд не влезают в 255 символов формулы — нужен диапазон."""
    ws = load_workbook(book_path)["Конструктор меню"]
    lists = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    recipe_rules = [dv for dv in lists if RECIPE_REF_SHEET in str(dv.formula1)]
    assert recipe_rules, "проверка ввода блюда должна ссылаться на справочник"
    assert len(recipe_rules[0].formula1) < 255


# ---------------------------------------------------------------------------
# Внутренние ссылки
# ---------------------------------------------------------------------------


def _links(ws) -> dict[str, str]:
    return {
        str(ws.cell(row=r, column=col).value): ws.cell(row=r, column=col).hyperlink.location
        for r in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
        if ws.cell(row=r, column=col).hyperlink
    }


def test_из_ттк_блюдо_ведёт_на_состав(book_path: Path):
    wb = load_workbook(book_path)
    links = _links(wb["ТТК"])
    assert set(links) == set(CATALOG_RECIPES)
    for name, location in links.items():
        assert location.startswith("'Ингредиенты'!B"), f"{name}: {location}"
        target = int(location.split("!B")[1])
        assert wb["Ингредиенты"].cell(row=target, column=2).value == name, (
            f"{name}: ссылка ведёт не на своё блюдо"
        )


def test_из_меню_блюдо_ведёт_на_состав(book_path: Path):
    wb = load_workbook(book_path)
    links = _links(wb["Меню"])
    assert set(links) == {m.recipe for m in demo_data(DAY).menu}
    for location in links.values():
        assert int(location.split("!B")[1]) > 0


def test_из_конструктора_блюдо_ведёт_на_состав(book_path: Path):
    wb = load_workbook(book_path)
    links = _links(wb["Конструктор меню"])
    assert set(links) == {l.recipe for l in demo_data(DAY).menu}
    for name, location in links.items():
        target = int(location.split("!B")[1])
        assert wb["Ингредиенты"].cell(row=target, column=2).value == name


def test_из_плана_блюдо_ведёт_на_состав(book_path: Path):
    wb = load_workbook(book_path)
    links = _links(wb["План меню"])
    assert set(links) == {p.recipe for p in WEEK_PLAN}
    for name, location in links.items():
        target = int(location.split("!B")[1])
        assert wb["Ингредиенты"].cell(row=target, column=2).value == name


def test_из_состава_есть_обратная_ссылка(book_path: Path):
    """Главная боль — потеряться в сотне строк и не найти свою карточку."""
    wb = load_workbook(book_path)
    ws = wb["Ингредиенты"]
    back = {
        str(ws.cell(row=r, column=2).value): ws.cell(row=r, column=8).hyperlink.location
        for r in range(1, ws.max_row + 1)
        if ws.cell(row=r, column=8).hyperlink
    }
    assert set(back) == set(CATALOG_RECIPES)
    for name, location in back.items():
        assert location.startswith("'ТТК'!B")
        target = int(location.split("!B")[1])
        assert wb["ТТК"].cell(row=target, column=2).value == name


def test_в_ингредиентах_название_блюда_один_раз_на_блок(book_path: Path):
    """Группировка: не «список строк», а блок на блюдо."""
    ws = load_workbook(book_path)["Ингредиенты"]
    header = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Блюдо"
    )
    named = [
        ws.cell(row=r, column=2).value
        for r in range(header + 1, ws.max_row + 1)
        if ws.cell(row=r, column=2).value
    ]
    assert len(named) == len(set(named)) == len(CATALOG_RECIPES)


def test_ссылки_не_ломают_пересборку(tmp_path: Path):
    """Файл с гиперссылками должен открываться и пересобираться без ошибок."""
    data = demo_data(DAY)
    path = tmp_path / "дважды.xlsx"
    build_workbook(data, build_shift_sheet(data), path)
    build_workbook(data, build_shift_sheet(data), path)
    wb = load_workbook(path)
    ttk = wb["ТТК"]
    header, _ = header_of(ttk, ("Блюдо", "Приём пищи"))
    dish = column_of(ttk, header, "Блюдо")
    first = row_of(ttk, dish, "Гороховый суп с копчёностями")
    assert ttk.cell(row=first, column=dish).hyperlink is not None


# ---------------------------------------------------------------------------
# Смена за день недели
# ---------------------------------------------------------------------------


def test_смена_за_день_считается_до_конца(book_path: Path):
    """Главный вопрос: --week даёт не список, а работающую кухню."""
    data = demo_data(DAY)
    menu = plan_for(read_plan(book_path, CATALOG_RECIPES), "пн")
    sheet = build_shift_sheet(replace(data, menu=menu, plan=()))

    assert len(sheet.menu) == len(plan_for(WEEK_PLAN, "пн"))
    assert sheet.tasks, "у дня недели должны быть задачи"
    assert sheet.needs_flat, "должна получиться продуктовая матрица"


def test_все_пять_дней_считаются_без_ошибок():
    data = demo_data(DAY)
    for day in WEEKDAYS:
        menu = plan_for(WEEK_PLAN, day)
        sheet = build_shift_sheet(replace(data, menu=menu, plan=()))
        assert sheet.tasks, f"{day}: задач нет"
        assert sheet.shift_end >= sheet.shift_start, f"{day}: смена вывернулась"


def test_план_в_книге_и_в_коде_совпадают(book_path: Path):
    """Шеф правит книгу, но пока он этого не сделал, план должен совпадать."""
    from_book = {p.weekday: set() for p in WEEK_PLAN}
    for line in read_plan(book_path, CATALOG_RECIPES):
        from_book[line.weekday].add(line.recipe)
    from_code = {p.weekday: set() for p in WEEK_PLAN}
    for line in WEEK_PLAN:
        from_code[line.weekday].add(line.recipe)
    assert from_book == from_code


def test_план_читается_даже_если_день_не_указан(book_path: Path):
    """День недели пишется один раз на блок — как в самой книге."""
    assert all(isinstance(p, PlanLine) for p in read_plan(book_path, CATALOG_RECIPES))


# ---------------------------------------------------------------------------
# Колонка «В меню» и предупреждения
# ---------------------------------------------------------------------------


def test_в_меню_отражает_смену_а_не_каталог(book_path: Path):
    """29 карточек в каталоге, но в работе сегодня — пять."""
    ws = load_workbook(book_path)["ТТК"]
    header, _ = header_of(ws, ("Блюдо", "В меню"))
    dish = column_of(ws, header, "Блюдо")
    flag = column_of(ws, header, "В меню")
    flags = {
        ws.cell(row=r, column=dish).value: ws.cell(row=r, column=flag).value
        for r in range(header + 1, ws.max_row + 1)
        if ws.cell(row=r, column=dish).value
    }
    in_menu = {l.recipe for l in demo_data(DAY).menu}
    assert set(flags) == set(CATALOG_RECIPES)
    assert {n for n, f in flags.items() if f == "да"} == in_menu
    assert sum(1 for f in flags.values() if f == "нет") == 24


def test_смена_за_день_недели_меняет_колонку_в_меню(tmp_path: Path):
    """В понедельник «да» у шести блюд, а не у пяти демо-блюд."""
    data = replace(demo_data(DAY), menu=plan_for(WEEK_PLAN, "пн"), plan=())
    path = build_workbook(data, build_shift_sheet(data), tmp_path / "пн.xlsx")
    ws = load_workbook(path)["ТТК"]
    header, _ = header_of(ws, ("Блюдо", "В меню"))
    flag = column_of(ws, header, "В меню")
    dish = column_of(ws, header, "Блюдо")
    yes = sum(
        1
        for r in range(header + 1, ws.max_row + 1)
        if ws.cell(row=r, column=flag).value == "да"
        and ws.cell(row=r, column=dish).value
    )
    assert yes == len(plan_for(WEEK_PLAN, "пн")) == 6


def test_блюдо_без_операций_даёт_предупреждение():
    """Карточка без строк в «Задачах» — тихая дырка в смене, если не сказать."""
    data = demo_data(DAY)
    recipes = dict(data.recipes)
    recipes["Тестовая запеканка"] = Recipe(
        name="Тестовая запеканка", meal="Обед", category="Бакалея",
        subcategory="Горячее", shop="Повар", prep_min=5, cook_min=20,
        ingredients=(Ingredient("Молоко", 50, "влить"),),
    )
    menu = data.menu + (
        MenuLine("Тестовая запеканка", 50, time(12, 0), "Обед", DAY),
    )
    sheet = build_shift_sheet(replace(data, recipes=recipes, menu=menu))
    assert any("Тестовая запеканка" in w for w in sheet.warnings)


def test_у_всех_блюд_плана_есть_цепочка_задач():
    """Обратная сторона: предупреждение не должно срабатывать на плане."""
    for day in WEEKDAYS:
        menu = plan_for(WEEK_PLAN, day)
        sheet = build_shift_sheet(replace(demo_data(DAY), menu=menu, plan=()))
        assert not any("Нет ни одной операции" in w for w in sheet.warnings), day
