"""Проверки слоя Excel: книга собирается, меню читается, расчёт пересчитывается.

Это самый ценный тест для пользователя: он повторяет его рабочий сценарий —
открыл файл, поменял порции, убрал блюдо, получил пересчитанную кухню.
"""

from __future__ import annotations

import re
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from kitchen.core import build_shift_sheet
from kitchen.excel import build_workbook, iter_blocks, read_menu
from kitchen.models import MenuLine
from kitchen.seed import demo_data

DAY = date(2026, 9, 26)

SHEETS = [
    "Лист смены",
    "Меню",
    "Продукты",
    "Карта задач",
    "Тайминг",
    "ТТК",
    "Ингредиенты",
    "Конструктор меню",
    "План меню",
    "Задачи",
    "Списания",
    "Инвентаризация",
    "Инструкция",
    "Справочник блюд",
]


@pytest.fixture()
def book_path(tmp_path: Path) -> Path:
    data = demo_data(DAY)
    path = tmp_path / "Кухня.xlsx"
    return build_workbook(data, build_shift_sheet(data), path)


def menu_of(path: Path) -> tuple[MenuLine, ...]:
    """Чтение меню из собранной книги."""
    return read_menu(path, demo_data(DAY).recipes)


def rewrite_menu(path: Path, *, portions: int | None = None, drop: str = "") -> None:
    """Правка листа «Конструктор меню» руками, как это делает пользователь."""
    wb = load_workbook(path)
    ws = wb["Конструктор меню"]
    for block in iter_blocks(ws):
        if block.dish is None or block.portions is None:
            continue
        for row in range(block.first, block.last + 1):
            name = ws.cell(row=row, column=block.dish).value
            if not name:
                continue
            if name == drop:
                ws.cell(row=row, column=block.portions, value=0)
            elif portions is not None:
                ws.cell(row=row, column=block.portions, value=portions)
    wb.save(path)


def free_slot(path: Path, meal: str) -> int:
    """Первая пустая строка окна приёма пищи — куда шеф ставит блюдо."""
    ws = load_workbook(path)["Конструктор меню"]
    for block in iter_blocks(ws):
        if block.meal != meal or block.dish is None:
            continue
        for row in range(block.first, block.last + 1):
            if not ws.cell(row=row, column=block.dish).value:
                return row
    raise AssertionError(f"в окне «{meal}» нет свободных строк")


def recompute(path: Path) -> object:
    """Полный цикл: книга → меню → расчёт, как в build.py."""
    data = demo_data(DAY)
    menu = read_menu(path, data.recipes)
    return build_shift_sheet(replace(data, menu=menu))


def test_книга_содержит_все_листы(book_path: Path):
    wb = load_workbook(book_path)
    assert wb.sheetnames == SHEETS


def test_лист_смены_открывается_первым(book_path: Path):
    assert load_workbook(book_path).sheetnames[0] == "Лист смены"


def test_конструктор_меню_читается_обратно(book_path: Path):
    menu = menu_of(book_path)
    assert len(menu) == 5
    assert {m.recipe for m in menu} == {l.recipe for l in demo_data(DAY).menu}
    assert all(m.portions == 50 for m in menu)


def test_правка_меню_пересчитывает_всю_кухню(book_path: Path):
    """60 порций вместо 50 и минус одно блюдо → другие продукты и задачи."""
    rewrite_menu(
        book_path,
        portions=60,
        drop="Салат из квашеной капусты и зелёного горошка",
    )
    sheet = recompute(book_path)

    assert len(sheet.menu) == 4, "блюдо с 0 порций не должно попадать в меню"
    assert all(r.portions == 60 for r in sheet.menu)

    # пропорционально выросло то, что зависит от порций
    farsh = next(n for n in sheet.needs_flat if n.product == "Фарш (говядина/свинина)")
    assert farsh.qty_display == pytest.approx(6.0 * 60 / 50)

    # горошек был только в салате: в меню он больше не нужен (0 г),
    # но остаётся в группе «проверить остаток» — вдруг на складе есть
    pea = next(n for n in sheet.needs_flat if n.product == "Горошек консервированный")
    assert pea.qty_display == 0
    grocery = next(g for g in sheet.needs if g.category == "Бакалея")
    assert pea in grocery.items and grocery.is_check_stock

    # число задач сократилось
    assert len(sheet.tasks) < len(build_shift_sheet(demo_data(DAY)).tasks)


def test_одна_строка_меню_не_ломает_остальные(book_path: Path):
    """Опечатка в одной строке не должна ронять весь расчёт."""
    wb = load_workbook(book_path)
    ws = wb["Конструктор меню"]
    row = free_slot(book_path, "Обед")
    block = next(b for b in iter_blocks(ws) if b.dish is not None and row in range(b.first, b.last + 1))
    ws.cell(row=row, column=block.dish, value="Борщ украинский")
    ws.cell(row=row, column=block.portions, value=30)
    wb.save(book_path)

    sheet = recompute(book_path)
    assert len(sheet.menu) == 5, "пять нормальных блюд должны уцелеть"
    assert all(r.dish != "Борщ украинский" for r in sheet.menu)
    assert all(r.portions == 50 for r in sheet.menu)


def test_пустой_файл_меню_не_ломает_сборку(tmp_path: Path):
    """Если пользователь стёр всё меню, книра всё равно должна собраться."""
    path = tmp_path / "пусто.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Конструктор меню"
    ws.append(["Блюдо", "Порций", "Выдача", "Приём пищи"])
    wb.save(path)

    assert read_menu(path, demo_data(DAY).recipes) == ()


def test_лист_меню_без_заголовка_даёт_пустое_меню(tmp_path: Path):
    """Файл, где заголовок случайно стёрли, — не повод падать."""
    path = tmp_path / "без_заголовка.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Конструктор меню"
    ws.append(["что-то", "не то"])
    wb.save(path)

    assert read_menu(path, demo_data(DAY).recipes) == ()


def test_файл_перезаписывается_без_ошибок(tmp_path: Path):
    data = demo_data(DAY)
    path = tmp_path / "дважды.xlsx"
    for _ in range(2):
        build_workbook(data, build_shift_sheet(data), path)
    assert load_workbook(path).sheetnames == SHEETS


def test_расчёт_не_сделан_формулами(book_path: Path):
    """Расчёт полностью на Python: битых формул в книге быть не должно.

    Исключение одно — счётчики панели поиска в «Конструкторе меню». Они
    считают совпадения запроса с ТТК и ничего не рассчитывают; без них поиск
    не работал бы без запуска скрипта.
    """
    wb = load_workbook(book_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                assert ws.title == "Конструктор меню", f"{ws.title}!{cell.coordinate}"
                assert cell.value.startswith(("=COUNTIF",)), cell.value


def test_лист_сцены_разбит_на_блоки_по_ролям(book_path: Path):
    ws = load_workbook(book_path)["Лист смены"]
    text = "\n".join(
        str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1)
    )
    for role in ("Разнорабочий", "Повар", "Шеф-повар", "Клининг"):
        assert role in text, f"роль {role} потерялась"


def test_примечание_не_рассыпается_по_символам(book_path: Path):
    """Регрессия: строка примечаний не должна делиться на посимвольные строки."""
    ws = load_workbook(book_path)["Лист смены"]
    notes = [
        ws.cell(row=r, column=1).value
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=1).value, str)
        and re.match(r"^\d+\. ", ws.cell(row=r, column=1).value)
    ]
    assert notes, "блок примечаний не найден"
    assert all(len(n) > 20 for n in notes), f"подозрение на разбивку: {notes}"
