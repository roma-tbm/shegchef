"""Регрессионные тесты дефектов третьего ревью Stage 3 (V3-01 … V3-07).

Тесты описывают ожидаемое поведение независимо от текущей реализации и были
написаны до исправления, чтобы сначала воспроизвести дефект. Проверяются
сквозные сценарии (меню → расчёт → печать → хранилище), а не отдельные
функции: почти каждый дефект проявляется только на стыке слоёв.

    V3-01  обязательная операция без исполнителя не снимает готовность
    V3-02  подготовка выходит за собственное окно
    V3-03  меню конкретного следующего дня не подключено в приложении
    V3-04  печать пустого субботнего отчёта падает
    V3-05  сохранение неполного отчёта удаляет предыдущий корректный
    V3-06  экспорт не использует сохранённый субботний факт
    V3-07  Excel и старые вкладки остаются отдельным рабочим планом
"""

from __future__ import annotations

from datetime import date, time
from pathlib import Path

import pytest

from kitchen.duties import DUTIES_BY_KEY, DutySettings
from kitchen.models import (
    Ingredient,
    KitchenData,
    MenuLine,
    Product,
    Recipe,
    TaskTemplate,
)
from kitchen.staff import ROLE_CHEF, ROLE_COOK, ROLE_HELPER, Employee
from kitchen.web.shift_plan import (
    PAGE_WIDTH,
    Printable,
    build_scenario,
    saturday_report_form,
)

MONDAY = date(2026, 10, 5)
TUESDAY = date(2026, 10, 6)
FRIDAY = date(2026, 10, 9)
SATURDAY = date(2026, 10, 10)
SUNDAY = date(2026, 10, 4)

#: Реальное блюдо из демо-набора: расчёт падает на несуществующей карточке.
NEXT_DAY_DISH = "Каша овсяная с солёным арахисом"

POTATO = Product(name="Картофель", unit="кг", category="Овощи", storage="Овощная")

STAGE_COOK = "Приготовление"
STAGE_PREP = "Заготовка"


def _recipe(name: str, *, prep_min: int = 0, cook_min: int = 10, meal: str = "Ужин") -> Recipe:
    return Recipe(
        name=name,
        meal=meal,
        category=meal,
        subcategory="Второе",
        shop="Горячий цех",
        prep_min=prep_min,
        cook_min=cook_min,
        ingredients=(Ingredient("Картофель", 200.0),),
    )


def _cook_template(name: str, *, operation: str | None = None, minutes: int = 10) -> TaskTemplate:
    return TaskTemplate(
        recipe=name,
        order=1,
        stage=STAGE_COOK,
        role=ROLE_COOK,
        operation=operation or f"Cook {name}",
        duration_min=minutes,
        products="Картофель",
        meal="Ужин",
    )


def _prep_template(name: str, *, operation: str | None = None, minutes: int = 30) -> TaskTemplate:
    return TaskTemplate(
        recipe=name,
        order=1,
        stage=STAGE_PREP,
        role=ROLE_COOK,
        operation=operation or f"Prep {name}",
        duration_min=minutes,
        products="Картофель",
        meal="Ужин",
    )


def _data(
    *,
    recipes: dict[str, Recipe] | None = None,
    tasks: tuple[TaskTemplate, ...] = (),
) -> KitchenData:
    return KitchenData(
        products={"Картофель": POTATO},
        recipes=recipes or {},
        tasks=tasks,
    )


def _employee(
    name: str,
    role: str,
    start: time = time(7, 0),
    end: time = time(20, 0),
    order: int = 1,
) -> Employee:
    return Employee(
        name=name,
        role=role,
        shift_start=start,
        shift_end=end,
        breaks=(),
        order=order,
    )


def _cook(name: str = "Cook", start: time = time(7, 0), order: int = 1) -> Employee:
    return _employee(name, ROLE_COOK, start, time(20, 0), order)


def _line(
    name: str,
    day: date,
    *,
    portions: int = 20,
    serve_at: time = time(12, 0),
) -> MenuLine:
    return MenuLine(
        recipe=name,
        portions=portions,
        serve_at=serve_at,
        meal="Ужин",
        day=day,
    )


NO_DUTIES = DutySettings(disabled=frozenset(DUTIES_BY_KEY))
ONLY_NEXT_DAY_PREP = DutySettings(disabled=frozenset(DUTIES_BY_KEY) - {"next-day-prep"})


def _prep_only(minutes: int) -> tuple[KitchenData, tuple[MenuLine, ...], tuple[MenuLine, ...]]:
    """Меню понедельника готовится, меню вторника — заготавливается.

    Операции разнесены по стадиям, иначе обратный тайминг не сможет отделить
    работу текущего дня от подготовки следующего.
    """
    data = _data(
        recipes={"Today": _recipe("Today"), "Tomorrow": _recipe("Tomorrow", prep_min=minutes)},
        tasks=(_cook_template("Today"), _prep_template("Tomorrow", minutes=minutes)),
    )
    return data, (_line("Today", MONDAY),), (_line("Tomorrow", TUESDAY),)


# ---------------------------------------------------------------------------
# V3-01 — обязательная операция без исполнителя (AC03, AC06, AC15)
# ---------------------------------------------------------------------------


def _two_overlapping_cooks() -> tuple[KitchenData, tuple[MenuLine, ...]]:
    """Два блюда, по одной операции 10 минут, оба к 12:00, один повар."""
    names = ("Cook A", "Cook B")
    return (
        _data(
            recipes={n: _recipe(n) for n in names},
            tasks=tuple(_cook_template(n) for n in names),
        ),
        (_line("Cook A", MONDAY), _line("Cook B", MONDAY)),
    )


def test_operation_without_assignee_makes_the_plan_not_ready() -> None:
    """Работа без исполнителя — это конфликт ресурса, а не готовый план.

    План не может выглядеть законченным, если обязательная операция ТТК
    осталась без человека: маршрут её не напечатает, а шеф не увидит, что
    делать некому.
    """
    data, menu = _two_overlapping_cooks()

    scenario = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )

    cook_items = [i for i in scenario.plan.items if i.role == ROLE_COOK]
    assert len(cook_items) == 2, "обе операции должны попасть в план"
    orphans = [i for i in cook_items if not i.assignee]
    assert orphans, "один повар не может взять две накладные операции"
    assert scenario.plan.ready is False, "план без исполнителя не готов"


def test_no_assignee_gap_names_the_work_it_blocks() -> None:
    """Причина без исполнителя называет саму работу, а не только роль."""
    data, menu = _two_overlapping_cooks()

    scenario = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )
    orphan = next(
        i for i in scenario.plan.items if i.role == ROLE_COOK and not i.assignee
    )

    gap = next(g for g in scenario.plan.gaps if orphan.item_id in g.code)
    assert orphan.operation in gap.subject
    assert gap.blocks, "причина должна называть, что именно не получается"


def test_fully_assigned_plan_stays_ready() -> None:
    """С двумя поварами те же операции назначены, и план снова готов.

    Обратная проверка: без неё тест выше проходил бы и при запрете готовить
    вовсе — то есть при «готовом» плане без единой операции.
    """
    data, menu = _two_overlapping_cooks()

    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        staff=(_cook("Cook 1"), _cook("Cook 2", order=2)),
        duty_settings=NO_DUTIES,
    )

    assert all(i.assignee for i in scenario.plan.items)
    assert scenario.plan.ready is True, [g.message for g in scenario.plan.gaps]


def test_work_before_the_employee_starts_is_not_ready() -> None:
    """Работа до начала смены сотрудника не может быть готовым планом.

    Обратный тайминг посчитает 06:00–07:00, но повар начинает в 07:00.
    Назначить работу раньше прихода нельзя, и это должно быть видно.
    """
    data = _data(
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(_cook_template("Cook A", minutes=60),),
    )

    scenario = build_scenario(
        data,
        MONDAY,
        menu=(_line("Cook A", MONDAY, serve_at=time(7, 0)),),
        staff=(_cook(start=time(7, 0)),),
        duty_settings=NO_DUTIES,
    )

    item = next(i for i in scenario.plan.items if i.operation == "Cook Cook A")
    assert not item.assignee, "раньше прихода сотрудника назначать нельзя"
    assert scenario.plan.ready is False


def test_no_assignee_is_never_printed_as_a_route() -> None:
    """Работа без исполнителя не попадает ни в один маршрут.

    Иначе шеф берёт в руки маршрут человека, в котором есть работа не его.
    """
    data, menu = _two_overlapping_cooks()

    scenario = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )

    for route in scenario.plan.routes:
        assert all(i.assignee == route.employee.name for i in route.items)


# ---------------------------------------------------------------------------
# V3-02 — подготовка обязана помещаться в своё окно (AC06, AC11)
# ---------------------------------------------------------------------------


def _prep_scenario(minutes: int):
    data, menu, next_menu = _prep_only(minutes)
    return build_scenario(
        data,
        MONDAY,
        menu=menu,
        next_menu=next_menu,
        staff=(_cook(),),
        duty_settings=ONLY_NEXT_DAY_PREP,
    )


def test_next_day_prep_may_not_leave_its_window() -> None:
    """Подготовка 240 минут не помещается в окно 14:00–17:00.

    Выход за собственное окно обязан обнаруживаться независимо от того,
    хватило ли свободного времени: иначе смена выглядит готовым расписанием,
    в котором часть работ заведомо не может быть выполнена в срок.
    """
    scenario = _prep_scenario(240)

    prep = [i for i in scenario.plan.items if i.operation.startswith("Prep Tomorrow")]
    assert prep, "подготовка следующего дня должна быть в плане"
    assert any(i.violations for i in prep), [i.span for i in prep]
    assert scenario.plan.ready is False
    assert any(
        "окно" in v.rule.lower() for i in prep for v in i.violations
    ), [str(v) for i in prep for v in i.violations]


def test_next_day_prep_inside_its_window_is_not_blocked() -> None:
    """Подготовка короче окна остаётся без нарушений (обратная проверка)."""
    scenario = _prep_scenario(120)

    prep = next(
        i for i in scenario.plan.items if i.operation.startswith("Prep Tomorrow")
    )
    assert prep.violations == (), prep.violations
    assert prep.end.time() <= time(17, 0)


def test_prep_after_friday_cutoff_is_not_planned() -> None:
    """В пятницу подготовка следующего дня не планируется после 14:00.

    Пятничный запрет (AC08) имеет приоритет над обычным шаблоном подготовки:
    работа не сокращается молча и не переносится на вечер.
    """
    data = _data(
        recipes={"Today": _recipe("Today"), "Tomorrow": _recipe("Tomorrow")},
        tasks=(_cook_template("Today"), _prep_template("Tomorrow", minutes=240)),
    )

    scenario = build_scenario(
        data,
        FRIDAY,
        menu=(_line("Today", FRIDAY),),
        next_menu=(_line("Tomorrow", SATURDAY),),
        staff=(_cook(),),
        duty_settings=ONLY_NEXT_DAY_PREP,
    )

    prep = [i for i in scenario.plan.items if i.operation.startswith("Prep Tomorrow")]
    assert not prep or all(i.violations for i in prep), [
        (i.span, i.violations) for i in prep
    ]
    assert scenario.plan.ready is False


def test_prep_needs_applicable_day_rule() -> None:
    """Подготовка порождается только режимом, который её допускает.

    Если обязанность не действует в этом дне, её содержание и не должно
    выводиться из меню: план не должен содержать работу, которой в этом
    дне быть не может.
    """
    data, menu, next_menu = _prep_only(30)

    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        next_menu=next_menu,
        staff=(_cook(),),
        # Обязанность вычеркнута настройкой — режим её не допускает.
        duty_settings=NO_DUTIES,
    )

    assert not [i for i in scenario.plan.items if i.operation.startswith("Prep ")]


def test_missing_next_menu_keeps_the_reason_visible() -> None:
    """Нет меню следующего дня — причина остаётся, а не исчезает молча."""
    data, menu, _ = _prep_only(30)

    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        next_menu=(),
        staff=(_cook(),),
        duty_settings=ONLY_NEXT_DAY_PREP,
    )

    assert any(
        "меню следующего дня" in (g.reason or "").lower() for g in scenario.plan.gaps
    ), [g.message for g in scenario.plan.gaps]
    assert scenario.plan.ready is False


# ---------------------------------------------------------------------------
# V3-03 — меню конкретного следующего дня (AC10, AC11)
# ---------------------------------------------------------------------------


def test_next_day_menu_is_the_menu_of_the_specific_day() -> None:
    """Подготовка берётся из меню нужного дня, а не ближайшего заполненного.

    Молчаливая подмена превращает воскресный план в план вторника: увидели
    «меню следующего дня» и нарезали не то.
    """
    from kitchen.web.shift_plan import next_day_menu

    data = _data(
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(),
    ).__class__(
        products={"Картофель": POTATO},
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(),
        menu=(_line("Cook A", TUESDAY),),
    )

    # Вторник заполнен, понедельник пуст — именно этот случай раньше молча
    # подставлял меню вторника.
    assert next_day_menu(data, SUNDAY) == ()
    assert next_day_menu(data, SUNDAY) != tuple(
        line for line in data.menu if line.day == TUESDAY
    )


def test_next_menu_lookup_uses_the_nearest_monday_for_sunday() -> None:
    """Воскресенье смотрит именно на ближайший понедельник."""
    from kitchen.web.shift_plan import next_day_menu, prep_target_day

    assert prep_target_day(SUNDAY) == MONDAY
    assert prep_target_day(MONDAY) == TUESDAY
    assert prep_target_day(FRIDAY) == SATURDAY

    data = KitchenData(
        products={"Картофель": POTATO},
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(),
        menu=(_line("Cook A", MONDAY), _line("Cook A", TUESDAY)),
    )

    from_monday = next_day_menu(data, SUNDAY)
    assert [line.day for line in from_monday] == [MONDAY]


def test_sunday_prep_uses_the_monday_menu_end_to_end() -> None:
    """Воскресный план строится из меню понедельника, порций и ТТК."""
    data = _data(
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(_prep_template("Cook A", minutes=30),),
    )
    monday_menu = (_line("Cook A", MONDAY, portions=30),)

    scenario = build_scenario(
        data,
        SUNDAY,
        menu=(),
        next_menu=monday_menu,
        staff=(_cook(),),
        duty_settings=DutySettings(disabled=frozenset(DUTIES_BY_KEY) - {"sunday-prep"}),
    )

    prep = next(
        (i for i in scenario.plan.items if "Cook A" in i.operation), None
    )
    assert prep is not None, "подготовка по меню понедельника должна появиться"
    assert prep.portions == 30


def test_sunday_without_monday_menu_has_no_invented_prep() -> None:
    """Нет меню понедельника — причина, а не заготовка чужого дня."""
    data = _data(
        recipes={"Cook A": _recipe("Cook A"), "Cook B": _recipe("Cook B")},
        tasks=(_prep_template("Cook A", minutes=30), _prep_template("Cook B", minutes=30)),
    )

    scenario = build_scenario(
        data,
        SUNDAY,
        menu=(),
        # Во вторник блюдо есть, в понедельник нет: подменять нельзя.
        next_menu=(),
        staff=(_cook(),),
        duty_settings=DutySettings(disabled=frozenset(DUTIES_BY_KEY) - {"sunday-prep"}),
    )
    assert scenario.next_menu == ()
    assert not [i for i in scenario.plan.items if "Cook B" in i.operation]
    assert scenario.plan.ready is False


def test_app_passes_the_user_menu_of_the_next_day(tmp_path, monkeypatch) -> None:
    """Streamlit доносит сохранённое меню следующего дня до расчёта.

    Сценарий сквозной: пользователь заполняет понедельник, открывает
    воскресенье — подготовка должна опираться на это меню.
    """
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    from kitchen.web import menu_store
    from kitchen.web.dataset import Dataset

    app_path = Path(__file__).resolve().parent.parent / "app.py"
    monkeypatch.setenv("SHEGCHEF_DATA_DIR", str(tmp_path))
    store = Dataset.in_dir(tmp_path).menu_alt
    menu_store.save_days(
        {
            MONDAY.isoformat(): [_line(NEXT_DAY_DISH, MONDAY)],
            SUNDAY.isoformat(): [],
        },
        store,
    )

    tester = AppTest.from_file(app_path, default_timeout=120).run()
    assert not tester.exception, [e.value for e in tester.exception]
    tester.session_state["day"] = SUNDAY
    tester.session_state["menu"] = ()
    tester.run()

    assert not tester.exception, [e.value for e in tester.exception]
    assert [
        line.recipe for line in tester.session_state["scenario"].next_menu
    ] == [NEXT_DAY_DISH], "меню понедельника должно доходить до воскресного расчёта"


def test_cli_passes_the_user_menu_of_the_next_day(tmp_path) -> None:
    """CLI доносит сохранённое меню следующего дня до расчёта."""
    import build

    from kitchen.web import menu_store

    store = tmp_path / "меню.json"
    menu_store.save_days(
        {
            MONDAY.isoformat(): [_line(NEXT_DAY_DISH, MONDAY)],
            SUNDAY.isoformat(): [],
        },
        store,
    )

    scenario = build._scenario_for(SUNDAY, store)

    assert [line.recipe for line in scenario.next_menu] == [NEXT_DAY_DISH]


# ---------------------------------------------------------------------------
# V3-04 — печать пустого субботнего отчёта (AC09, AC14)
# ---------------------------------------------------------------------------


def test_empty_saturday_report_prints_without_explosion() -> None:
    """Пустой отчёт за субботу печатается, а не роняет экспорт."""
    scenario = build_scenario(KitchenData(), SATURDAY, menu=())

    text = saturday_report_form(scenario).to_text()

    form = saturday_report_form(scenario)

    assert text.strip(), "пустая форма всё равно должна печататься"
    assert form.min_page_width() <= PAGE_WIDTH, (
        "пустая форма должна помещаться на лист, иначе её не напечатать"
    )
    for line in text.splitlines():
        assert len(line) <= PAGE_WIDTH, f"строка длиннее листа: {line!r}"


def test_empty_saturday_report_rows_match_the_header() -> None:
    """Каждая строка отчёта содержит столько же ячеек, сколько заголовок."""
    form = saturday_report_form(build_scenario(KitchenData(), SATURDAY, menu=()))

    assert form.rows, "пустая форма всё равно содержит строку-заглушку"
    for row in form.rows:
        assert len(row) == len(form.header)


def test_empty_saturday_report_names_the_missing_facts() -> None:
    """Пустая форма говорит, что расход и маркировка не заполнены."""
    form = saturday_report_form(build_scenario(KitchenData(), SATURDAY, menu=()))

    notes = " ".join(form.notes).lower()
    assert "расход" in notes
    assert "выход" in notes
    assert "маркиров" in notes


def test_empty_saturday_report_does_not_invent_a_zero() -> None:
    """Отсутствующий факт печатается прочерком, а не нулём."""
    form = saturday_report_form(build_scenario(KitchenData(), SATURDAY, menu=()))

    assert form.rows[0][3] == "—"
    assert form.rows[0][5] == "—"


# ---------------------------------------------------------------------------
# V3-05 — неполный отчёт не удаляет предыдущий корректный (AC09)
# ---------------------------------------------------------------------------


def _complete_report(day: date):
    from kitchen.web.report_store import (
        SaturdayReport,
        SemifinishedMarking,
        SemifinishedOutput,
    )

    return SaturdayReport(
        day=day,
        consumption={"Картофель": 5.0},
        outputs={"Бульон говяжий": SemifinishedOutput(qty=4.0, unit="кг")},
        markings={
            "Бульон говяжий": SemifinishedMarking(
                name="Бульон говяжий", made=day, expires=date(2026, 10, 13)
            )
        },
    )


def _incomplete_report(day: date):
    from kitchen.web.report_store import SaturdayReport, SemifinishedOutput

    return SaturdayReport(
        day=day,
        consumption={},
        outputs={"Бульон говяжий": SemifinishedOutput(qty=4.0, unit="кг")},
    )


def test_incomplete_report_raises_validation_error(tmp_path) -> None:
    """Неполный отчёт — ошибка валидации, а не тихая потеря записи."""
    from kitchen.web.report_store import IncompleteReportError, save_report

    with pytest.raises(IncompleteReportError):
        save_report(SATURDAY, _incomplete_report(SATURDAY), tmp_path / "reports.json")


def test_incomplete_report_keeps_the_previous_correct_file(tmp_path) -> None:
    """Предыдущий корректный отчёт остаётся в файле без изменений."""
    from kitchen.web import report_store

    path = tmp_path / "reports.json"
    report_store.save_report(SATURDAY, _complete_report(SATURDAY), path)
    before = path.read_text(encoding="utf-8")

    with pytest.raises(report_store.IncompleteReportError):
        report_store.save_report(SATURDAY, _incomplete_report(SATURDAY), path)

    assert path.read_text(encoding="utf-8") == before
    loaded = report_store.load_report(SATURDAY, path)
    assert loaded is not None and loaded.consumption == {"Картофель": 5.0}


def test_incomplete_report_does_not_touch_other_days(tmp_path) -> None:
    """Ошибка одного дня не трогает отчёты соседних суббот."""
    from kitchen.web import report_store

    path = tmp_path / "reports.json"
    other = date(2026, 9, 26)
    report_store.save_report(other, _complete_report(other), path)

    with pytest.raises(report_store.IncompleteReportError):
        report_store.save_report(SATURDAY, _incomplete_report(SATURDAY), path)

    assert report_store.load_report(other, path) is not None
    assert report_store.load_report(SATURDAY, path) is None


def test_incomplete_report_error_names_the_semifinished_product(tmp_path) -> None:
    """Ошибка называет полуфабрикат без маркировки."""
    from kitchen.web.report_store import IncompleteReportError, save_report

    with pytest.raises(IncompleteReportError, match="Бульон говяжий"):
        save_report(SATURDAY, _incomplete_report(SATURDAY), tmp_path / "reports.json")


# ---------------------------------------------------------------------------
# V3-06 — экспорт использует сохранённый субботний факт (AC09, AC14)
# ---------------------------------------------------------------------------


def _saved_report(tmp_path):
    from kitchen.web import report_store

    path = tmp_path / "reports.json"
    report_store.save_report(SATURDAY, _complete_report(SATURDAY), path)
    return path


def test_export_uses_the_saved_saturday_facts_after_restart(tmp_path) -> None:
    """Главный экспорт отдаёт сохранённый отчёт, а не пустой шаблон.

    Сценарий сквозной: ввод → сохранение → «перезапуск» → экспорт. Иначе
    отчёт, который человек заполнил в субботу, выглядит потерянным уже на
    следующий день.
    """
    from kitchen.web import report_store

    path = _saved_report(tmp_path)
    scenario = build_scenario(KitchenData(), SATURDAY, menu=())
    saved = report_store.load_report(SATURDAY, path)

    text = saturday_report_form(scenario, **_facts(saved)).to_text()

    assert "Картофель" in text
    assert "5" in text
    assert "4" in text
    assert "кг" in text
    assert "13.10.2026" in text


def _facts(saved):
    if saved is None:
        return {}
    return {
        "actuals": saved.consumption,
        "outputs": saved.outputs,
        "markings": saved.markings,
    }


def test_export_without_saved_report_prints_an_empty_form(tmp_path) -> None:
    """Без сохранённого факта экспорт печатает пустую форму, а не падает."""
    from kitchen.web import report_store

    scenario = build_scenario(KitchenData(), SATURDAY, menu=())
    saved = report_store.load_report(SATURDAY, tmp_path / "reports.json")

    text = saturday_report_form(scenario, **_facts(saved)).to_text()

    assert text.strip()
    assert "Фактический расход не введён" in text


def test_cli_export_uses_the_saved_saturday_facts(tmp_path) -> None:
    """CLI-экспорт субботнего отчёта читает сохранённый факт."""
    import build

    from kitchen.web import report_store

    path = _saved_report(tmp_path)
    scenario = build_scenario(KitchenData(), SATURDAY, menu=())
    saved = report_store.load_report(SATURDAY, path)

    text = saturday_report_form(scenario, **_facts(saved)).to_text()
    assert "Картофель" in text
    # `_saved_report_for` возвращает весь SaturdayReport (его принимает
    # `_saved_facts`), поэтому сверяется загруженный факт, а не одно поле.
    loaded = build._saved_report_for(SATURDAY, path.parent)
    assert loaded is not None and loaded.consumption == saved.consumption


def test_cli_printout_uses_the_saved_saturday_facts(tmp_path) -> None:
    """Настоящий путь `build.main` печатает сохранённый факт, не пустой шаблон.

    Прямые вызовы печатной формы не проверяют, что CLI действительно читает
    файл отчётов рядом с книгой. Здесь запускается реальный main с --no-open.
    """
    import build as cli

    from kitchen.web import report_store

    out = tmp_path / "Кухня.xlsx"
    report_store.save_report(
        SATURDAY, _complete_report(SATURDAY), tmp_path / "reports.json"
    )

    assert cli.main(["--out", str(out), "--day", "10.10.2026", "--menu", "50",
                     "--reports", str(tmp_path / "reports.json"), "--no-open"]) == 0

    forms = list((tmp_path / cli.PRINT_DIR_NAME).glob("Субботний отчёт*.txt"))
    assert forms, "субботний отчёт должен быть выгружен"
    text = forms[0].read_text(encoding="utf-8")
    assert "Картофель" in text
    assert "Бульон говяжий" in text
    assert "13.10.2026" in text


def test_saved_report_survives_a_broken_file(tmp_path) -> None:
    """Повреждённый файл отчётов — понятная ошибка, а не выдуманный факт (AC06)."""
    import pytest

    from kitchen.web import report_store

    path = tmp_path / "reports.json"
    path.write_text("{не json", encoding="utf-8")

    with pytest.raises(report_store.ReportStoreError):
        report_store.load_report(SATURDAY, path)
    assert path.read_text(encoding="utf-8") == "{не json"


def test_report_of_a_missing_file_is_none(tmp_path) -> None:
    """Отсутствующего файла отчётов это не касается — это отдельное состояние."""
    from kitchen.web import report_store

    assert report_store.load_report(SATURDAY, tmp_path / "нет.json") is None


# ---------------------------------------------------------------------------
# V3-07 — один авторитетный план для плана, маршрутов и книги (AC12)
# ---------------------------------------------------------------------------


def _late_shift_scenario():
    """Обеденная операция 60 минут, повар приходит в 09:00 — позже выдачи."""
    data = _data(
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(_cook_template("Cook A", minutes=60),),
    )
    return build_scenario(
        data,
        MONDAY,
        menu=(_line("Cook A", MONDAY, serve_at=time(8, 0)),),
        staff=(_cook("Cook", start=time(9, 0)),),
        duty_settings=NO_DUTIES,
    )


def test_workbook_timeline_uses_the_planner_schedule() -> None:
    """Книга берёт время из авторитетного плана, а не из своего тайминга.

    Старое расписание начинает сутки с 06:00 и рисует 06:00–07:00 повару,
    который приходит в 09:00. Такая книга противоречит маршруту, который
    человек берёт в руки (AC12).
    """
    from kitchen.core import build_shift_sheet

    data = _data(
        recipes={"Cook A": _recipe("Cook A")},
        tasks=(_cook_template("Cook A", minutes=60),),
    )
    menu = (_line("Cook A", MONDAY, serve_at=time(8, 0)),)
    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        staff=(_cook("Cook", start=time(9, 0)),),
        duty_settings=NO_DUTIES,
    )

    sheet = build_shift_sheet(replace_menu(data, menu), plan=scenario.plan)

    planned = {
        (i.task_id, i.start, i.end) for i in scenario.plan.items if i.task_id
    }
    for line in sheet.tasks:
        assert (line.task_id, line.start, line.end) in planned, (
            f"{line.operation} {line.start:%H:%M} не совпадает с расчётом"
        )


def replace_menu(data: KitchenData, menu: tuple[MenuLine, ...]) -> KitchenData:
    from dataclasses import replace

    return replace(data, menu=menu)


def test_workbook_reports_the_same_readiness() -> None:
    """Книга показывает те же неготовые причины, что и общий план."""
    from kitchen.core import build_shift_sheet

    data, menu = _two_overlapping_cooks()
    scenario = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )

    sheet = build_shift_sheet(replace_menu(data, menu), plan=scenario.plan)

    assert scenario.plan.ready is False
    text = "\n".join(sheet.warnings)
    assert "исполнител" in text.lower(), text


def test_workbook_keeps_its_existing_format(tmp_path) -> None:
    """Формат книги прежний: те же листы и те же колонки."""
    from openpyxl import load_workbook

    from kitchen.core import build_shift_sheet
    from kitchen.excel.builder import build_workbook

    data, menu = _two_overlapping_cooks()
    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        staff=(_cook("Cook 1"), _cook("Cook 2", order=2)),
        duty_settings=NO_DUTIES,
    )

    path = tmp_path / "Кухня.xlsx"
    build_workbook(data, build_shift_sheet(replace_menu(data, menu), plan=scenario.plan), path)

    wb = load_workbook(path)
    for sheet in ("Лист смены", "Меню", "Продукты", "Карта задач", "Тайминг"):
        assert sheet in wb.sheetnames, wb.sheetnames


def test_legacy_tabs_read_the_same_plan(tmp_path) -> None:
    """Старые вкладки приложения показывают план, а не отдельный расчёт."""
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    app_path = Path(__file__).resolve().parent.parent / "app.py"

    tester = AppTest.from_file(app_path, default_timeout=120).run()
    assert not tester.exception, [e.value for e in tester.exception]

    scenario = tester.session_state["scenario"]
    sheet = tester.session_state["sheet"]
    assert sheet is not None

    planned = {i.item_id: (i.start, i.end) for i in scenario.plan.items}
    seen = {line.item_id for line in sheet.tasks}
    # Прямая и обратная полнота: ни одна работа плана не потеряна и ни одна
    # строка книги не взята из чужого расчёта.
    assert set(planned) <= seen, sorted(set(planned) - seen)
    for line in sheet.tasks:
        assert line.item_id in planned, line.operation
        assert (line.start, line.end) == planned[line.item_id], line.operation


def test_working_schedule_survives_a_second_recalculation() -> None:
    """Повторный расчёт неизменных данных даёт тот же план (AC13)."""
    data, menu = _two_overlapping_cooks()
    kwargs = dict(
        menu=menu,
        staff=(_cook("Cook 1"), _cook("Cook 2", order=2)),
        duty_settings=NO_DUTIES,
    )

    first = build_scenario(data, MONDAY, **kwargs)
    second = build_scenario(data, MONDAY, **kwargs)

    assert [(i.item_id, i.span, i.assignee) for i in first.plan.items] == [
        (i.item_id, i.span, i.assignee) for i in second.plan.items
    ]


def test_chef_and_helper_share_the_same_shift_plan() -> None:
    """Шеф, повар и разнорабочий видят одно расписание, а не каждый своё."""
    from kitchen.core import build_shift_sheet

    data, menu = _two_overlapping_cooks()
    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        staff=(
            _employee("Шеф", ROLE_CHEF, time(8, 0), time(17, 0), 1),
            _cook("Cook", order=2),
            _employee("Разнорабочий", ROLE_HELPER, time(8, 0), time(20, 0), 3),
        ),
        duty_settings=NO_DUTIES,
    )

    sheet = build_shift_sheet(replace_menu(data, menu), plan=scenario.plan)
    planned = {(i.task_id, i.start, i.end) for i in scenario.plan.items if i.task_id}

    assert all((t.task_id, t.start, t.end) in planned for t in sheet.tasks)
    assert sheet.tasks


# ---------------------------------------------------------------------------
# Сквозная честность готовности
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("cooks", (1, 2), ids=("one-cook", "two-cooks"))
def test_readiness_matches_the_work_that_is_actually_assigned(cooks: int) -> None:
    """Готовность честна при любом составе из одного или двух поваров."""
    data, menu = _two_overlapping_cooks()
    staff = tuple(_cook(f"Cook {i + 1}", order=i + 1) for i in range(cooks))

    scenario = build_scenario(data, MONDAY, menu=menu, staff=staff, duty_settings=NO_DUTIES)

    orphans = [i for i in scenario.plan.items if not i.assignee and i.start < i.end]
    assert scenario.plan.ready is (not orphans)


def test_unknown_roster_is_reported_once() -> None:
    """Незаданный состав выходного дня объясняется один раз, а не по пунктам."""
    scenario = build_scenario(
        KitchenData(),
        SATURDAY,
        menu=(),
        staff=(),
        duty_settings=DutySettings(),
    )

    reasons = [g.reason for g in scenario.plan.gaps]
    assert any("состав" in (r or "").lower() for r in reasons)
    assert scenario.plan.ready is False
