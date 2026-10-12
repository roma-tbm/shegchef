"""Регрессионные тесты дефектов пятого ревью Stage 3 (V5-01, V5-02).

    V5-01  книга Excel подменяла расчётный BLOCKED пользовательской отметкой DONE
    V5-02  контрольные точки не переносили статус, факт и комментарий

Тесты смотрят на реальную книгу openpyxl: заблокированная работа обязана
остаться ⛔/BLOCKED во всех рабочих представлениях, а заявленный человеком
факт печататься отдельно («факт: DONE»).
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook

from kitchen.core import build_shift_sheet
from kitchen.duties import DUTIES_BY_KEY, DutySettings
from kitchen.excel.builder import build_workbook
from kitchen.models import Ingredient, KitchenData, MenuLine, Product, Recipe, TaskTemplate
from kitchen.staff import ROLE_COOK, Employee
from kitchen.web.messages import messages_by_role
from kitchen.web.shift_plan import build_scenario, route_form

MONDAY = date(2026, 10, 5)
FRIDAY = date(2026, 10, 9)

POTATO = Product(name="Картофель", unit="кг", category="Овощи", storage="Овощная")
NO_DUTIES = DutySettings(disabled=frozenset(DUTIES_BY_KEY))


def _recipe(name: str) -> Recipe:
    return Recipe(
        name=name,
        meal="Обед",
        category="Обед",
        subcategory="Второе",
        shop="Горячий цех",
        prep_min=0,
        cook_min=10,
        ingredients=(Ingredient("Картофель", 200.0),),
    )


def _template(
    name: str, *, operation: str, stage: str = "Приготовление", minutes: int = 10
) -> TaskTemplate:
    return TaskTemplate(
        recipe=name,
        order=1,
        stage=stage,
        role=ROLE_COOK,
        operation=operation,
        duration_min=minutes,
        products="Картофель",
        meal="Обед",
    )


def _line(name: str, day: date, *, portions: int = 20, serve_at: time = time(12, 0)) -> MenuLine:
    return MenuLine(recipe=name, portions=portions, serve_at=serve_at, meal="Обед", day=day)


def _cook() -> Employee:
    return Employee(
        name="Cook", role=ROLE_COOK, shift_start=time(7, 0), shift_end=time(20, 0),
        breaks=(), order=1,
    )


def _menu_data(name: str, day: date, template: TaskTemplate) -> KitchenData:
    return KitchenData(
        products={"Картофель": POTATO},
        recipes={name: _recipe(name)},
        tasks=(template,),
        menu=(_line(name, day),),
    )


def _with_menu(data: KitchenData, menu: tuple[MenuLine, ...]) -> KitchenData:
    return replace(data, menu=menu)


def _rows_of(ws, needle: str) -> list[list[str]]:
    """Все строки листа, где встречается значение needle (list значений)."""
    out: list[list[str]] = []
    for row in ws.iter_rows():
        values = [cell.value for cell in row]
        if any(isinstance(v, str) and needle in v for v in values):
            out.append(["" if v is None else str(v) for v in values])
    assert out, f"{needle!r} не найдено на листе «{ws.title}»"
    return out


def _blocked_friday_scenario():
    """Пятничная операция после 14:00 — расчётно заблокированная работа."""
    data = _menu_data(
        "Late dish",
        FRIDAY,
        _template("Late dish", operation="Forbidden cook"),
    )
    menu = (_line("Late dish", FRIDAY, serve_at=time(15, 0)),)

    draft = build_scenario(
        data, FRIDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )
    item = next(i for i in draft.plan.items if i.operation == "Forbidden cook")
    assert item.status == "BLOCKED", item.status
    assert draft.plan.ready is False

    scenario = build_scenario(
        data,
        FRIDAY,
        menu=menu,
        staff=(_cook(),),
        duty_settings=NO_DUTIES,
        overrides=((item.item_id, "DONE"),),
    )
    blocked = next(i for i in scenario.plan.items if i.operation == "Forbidden cook")
    return data, menu, scenario, blocked


# ---------------------------------------------------------------------------
# V5-01 — BLOCKED не подменяется отметкой DONE
# ---------------------------------------------------------------------------


def test_claimed_done_does_not_replace_blocked_status() -> None:
    """Расчётный статус остаётся BLOCKED, факт DONE хранится отдельно."""
    _, _, scenario, blocked = _blocked_friday_scenario()

    assert blocked.status == "BLOCKED", "расчётная блокировка не должна меняться"
    assert blocked.claimed_status == "DONE"
    assert scenario.plan.ready is False, "отметка человека не снимает неготовность"


def test_workbook_keeps_blocked_in_task_map_and_timeline(tmp_path) -> None:
    """Книга показывает ⛔/BLOCKED, а «факт: DONE» — отдельно."""
    data, menu, scenario, _ = _blocked_friday_scenario()
    book_data = _with_menu(data, menu)
    sheet = build_shift_sheet(book_data, plan=scenario.plan)

    path = tmp_path / "Кухня.xlsx"
    build_workbook(book_data, sheet, path)
    wb = load_workbook(path)

    task_rows = _rows_of(wb["Карта задач"], "Forbidden cook")
    assert any("⛔" in r for r in task_rows), task_rows
    assert not any("☑" in r for r in task_rows), "DONE не должен подменять блокировку"
    assert any(any("факт: DONE" in v for v in r) for r in task_rows), task_rows

    timeline_rows = _rows_of(wb["Тайминг"], "Forbidden cook")
    assert any("BLOCKED" in r for r in timeline_rows), timeline_rows
    # Статус и факт стоят в разных колонках: «DONE» не подменяет «BLOCKED».
    assert any(
        "BLOCKED" in r and any(v == "факт: DONE" for v in r)
        for r in timeline_rows
    ), timeline_rows

    # «Лист смены»: та же работа остаётся заблокированной с отдельным фактом.
    shift_rows = _rows_of(wb["Лист смены"], "Forbidden cook")
    assert any("⛔" in r for r in shift_rows), shift_rows
    assert any(any("факт: DONE" in v for v in r) for r in shift_rows), shift_rows


def test_route_and_message_keep_blocked_with_claimed_fact() -> None:
    """Маршрут и текстовое задание тоже не прячут блокировку."""
    data, menu, scenario, _ = _blocked_friday_scenario()
    book_data = _with_menu(data, menu)
    sheet = build_shift_sheet(book_data, plan=scenario.plan)

    route = route_form(scenario, scenario.plan.staff[0]).to_text()
    assert "BLOCKED" in route
    assert "факт: DONE" in route

    messages = {m.role: m.text for m in messages_by_role(sheet)}
    text = messages["Повар"]
    assert "Статус: BLOCKED" in text
    assert "факт: DONE" in text


# ---------------------------------------------------------------------------
# V5-02 — контрольные точки несут статус, факт и комментарий
# ---------------------------------------------------------------------------


def _control_scenario(*, day: date = MONDAY, serve_at: time = time(12, 0)):
    data = _menu_data(
        "Control dish",
        day,
        _template(
            "Control dish", operation="Проверка блюда", stage="Контроль", minutes=10
        ),
    )
    data = _with_menu(data, (_line("Control dish", day, serve_at=serve_at),))
    return data


def test_control_point_keeps_status_fact_and_comment(tmp_path) -> None:
    """Контрольная точка с DONE и комментарием печатается в таблице контроля."""
    data = _control_scenario()
    menu = tuple(data.menu)

    draft = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )
    control = next(i for i in draft.plan.items if i.operation == "Проверка блюда")

    scenario = build_scenario(
        data,
        MONDAY,
        menu=menu,
        staff=(_cook(),),
        duty_settings=NO_DUTIES,
        overrides=((control.item_id, "DONE"),),
        notes=((control.item_id, "проверено шефом"),),
    )
    book_data = _with_menu(data, menu)
    sheet = build_shift_sheet(book_data, plan=scenario.plan)

    line = next(c for c in sheet.controls if c.operation == "Проверка блюда")
    assert line.status == "DONE"
    assert line.comment == "проверено шефом"
    assert line.item_id == control.item_id

    path = tmp_path / "Кухня.xlsx"
    build_workbook(book_data, sheet, path)
    wb = load_workbook(path)

    # «Проверка блюда» есть и в хронологии, и в таблице контроля: нужна
    # строка контроля, где печатается комментарий.
    rows = _rows_of(wb["Тайминг"], "Проверка блюда")
    assert any(
        any("проверено шефом" in v for v in r) for r in rows
    ), rows
    control_row = next(r for r in rows if any("проверено шефом" in v for v in r))
    assert "☑" in control_row, control_row


def test_blocked_control_point_keeps_claimed_fact(tmp_path) -> None:
    """Заблокированная контрольная точка остаётся ⛔ с отметкой «факт: DONE»."""
    data = _control_scenario()
    menu = tuple(data.menu)

    draft = build_scenario(
        data, MONDAY, menu=menu, staff=(_cook(),), duty_settings=NO_DUTIES
    )
    control = next(i for i in draft.plan.items if i.operation == "Проверка блюда")

    # Моделируем расчётно заблокированную точку с заявленным фактом: важно,
    # что отображение не подменяет статус отметкой человека.
    blocked_plan = replace(
        draft.plan,
        items=tuple(
            replace(
                i,
                status="BLOCKED",
                claimed_status="DONE",
                comment="шеф настоял",
                violations=(),
            )
            if i.item_id == control.item_id
            else i
            for i in draft.plan.items
        ),
    )
    sheet = build_shift_sheet(_with_menu(data, menu), plan=blocked_plan)

    line = next(c for c in sheet.controls if c.operation == "Проверка блюда")
    assert line.status == "BLOCKED"
    assert line.claimed_status == "DONE"

    path = tmp_path / "Кухня.xlsx"
    build_workbook(_with_menu(data, menu), sheet, path)
    wb = load_workbook(path)
    rows = _rows_of(wb["Тайминг"], "Проверка блюда")
    control_row = next(r for r in rows if any("шеф настоял" in v for v in r))
    assert "⛔" in control_row, control_row
    assert "☑" not in control_row, "отметка DONE не заменяет блокировку"
    assert any("факт: DONE" in v for v in control_row), control_row
